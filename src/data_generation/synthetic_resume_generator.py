"""
Synthetic Resume Generator
Generates realistic resumes with both genuine and fraudulent patterns
"""

import random
import json
from datetime import datetime, timedelta
from pathlib import Path
from typing import Dict, List, Tuple
import sys
sys.path.append(str(Path(__file__).parent.parent.parent))

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib import colors
from docx import Document
from docx.shared import Pt, Inches
from tqdm import tqdm
import config

# Sample data for resume generation
FIRST_NAMES = [
    "James", "John", "Robert", "Michael", "William", "David", "Richard", "Joseph",
    "Mary", "Patricia", "Jennifer", "Linda", "Elizabeth", "Barbara", "Susan", "Jessica",
    "Sarah", "Emily", "Emma", "Olivia", "Sophia", "Isabella", "Ava", "Mia"
]

LAST_NAMES = [
    "Smith", "Johnson", "Williams", "Brown", "Jones", "Garcia", "Miller", "Davis",
    "Rodriguez", "Martinez", "Hernandez", "Lopez", "Wilson", "Anderson", "Thomas",
    "Taylor", "Moore", "Jackson", "Martin", "Lee", "Thompson", "White", "Harris"
]

GENUINE_UNIVERSITIES = [
    "Stanford University", "MIT", "Harvard University", "UC Berkeley",
    "Carnegie Mellon University", "University of Michigan", "Georgia Tech",
    "University of Texas at Austin", "UCLA", "Columbia University",
    "Cornell University", "Princeton University", "Yale University"
]

FAKE_UNIVERSITIES = [
    "American International University", "Global Tech Institute",
    "International Business School", "Pacific Western University",
    "Trinity Southern University", "Universal Career Institute",
    "Worldwide Online University", "Elite Management Institute"
]

GENUINE_COMPANIES = [
    "Google", "Microsoft", "Amazon", "Apple", "Meta", "IBM", "Oracle",
    "Salesforce", "Adobe", "Intel", "Cisco", "Netflix", "Tesla",
    "Accenture", "Deloitte", "PWC", "Goldman Sachs", "JPMorgan Chase"
]

FAKE_COMPANIES = [
    "Global Tech Solutions Inc", "Advanced Systems Corp", "Premier Consulting Group",
    "International Business Partners", "Elite Technologies Ltd", "Worldwide Innovations",
    "Strategic Solutions International", "NextGen Enterprises"
]

DEGREES = {
    "Bachelor": ["Bachelor of Science in Computer Science", "Bachelor of Engineering",
                 "Bachelor of Technology", "Bachelor of Business Administration"],
    "Master": ["Master of Science in Computer Science", "Master of Business Administration",
               "Master of Engineering", "Master of Technology"],
    "PhD": ["PhD in Computer Science", "PhD in Engineering", "PhD in Business"]
}

JOB_TITLES_PROGRESSION = {
    "entry": ["Software Engineer", "Junior Developer", "Associate Consultant", "Analyst"],
    "mid": ["Senior Software Engineer", "Team Lead", "Manager", "Senior Consultant"],
    "senior": ["Principal Engineer", "Director", "Senior Manager", "VP Engineering"],
    "executive": ["CTO", "VP", "Chief Architect", "Head of Engineering"]
}

INFLATED_TITLES = [
    "Chief Innovation Officer", "Principal Architect and VP", "Executive Director",
    "Global Head of Technology", "Senior VP and Chief Strategist"
]

SKILLS = {
    "programming": ["Python", "Java", "C++", "JavaScript", "Go", "Rust", "TypeScript"],
    "web": ["React", "Angular", "Vue.js", "Node.js", "Django", "Flask", "Spring Boot"],
    "data": ["SQL", "MongoDB", "PostgreSQL", "Redis", "Elasticsearch"],
    "cloud": ["AWS", "Azure", "Google Cloud", "Docker", "Kubernetes"],
    "ml": ["TensorFlow", "PyTorch", "Scikit-learn", "Machine Learning", "Deep Learning"],
    "tools": ["Git", "Jenkins", "CI/CD", "Agile", "Scrum"]
}

UNREALISTIC_SKILL_COMBOS = [
    # Expert in too many unrelated fields
    ["Expert in 15+ Programming Languages", "Master of All Cloud Platforms",
     "Advanced AI/ML Guru", "Blockchain Expert", "Quantum Computing Specialist"],
]

CERTIFICATIONS_GENUINE = [
    "AWS Certified Solutions Architect", "Google Cloud Professional",
    "Microsoft Certified Azure Administrator", "PMP Certification",
    "Certified Scrum Master", "CISSP"
]

CERTIFICATIONS_FAKE = [
    "International IT Excellence Certificate", "Global Tech Mastery Award",
    "Advanced Systems Certification (Online)", "Premier Developer Certificate"
]


class ResumeGenerator:
    """Generates synthetic resumes with realistic or fraudulent patterns"""
    
    def __init__(self, output_dir: Path):
        self.output_dir = output_dir
        self.output_dir.mkdir(parents=True, exist_ok=True)
        self.metadata = []
        
    def generate_name(self) -> str:
        """Generate random person name"""
        return f"{random.choice(FIRST_NAMES)} {random.choice(LAST_NAMES)}"
    
    def generate_contact(self) -> Dict:
        """Generate contact information"""
        name = self.generate_name()
        email = f"{name.lower().replace(' ', '.')}@{'gmail.com' if random.random() > 0.3 else 'email.com'}"
        phone = f"+1-{random.randint(200, 999)}-{random.randint(200, 999)}-{random.randint(1000, 9999)}"
        return {
            "name": name,
            "email": email,
            "phone": phone,
            "address": f"{random.randint(100, 9999)} Main Street, City, State {random.randint(10000, 99999)}"
        }
    
    def generate_education(self, is_fraudulent: bool) -> List[Dict]:
        """Generate education history"""
        education = []
        current_year = datetime.now().year
        
        # Bachelor's degree
        bachelors_end = current_year - random.randint(2, 15)
        bachelors_start = bachelors_end - 4
        
        university = random.choice(FAKE_UNIVERSITIES if is_fraudulent and random.random() > 0.5 
                                  else GENUINE_UNIVERSITIES)
        
        education.append({
            "degree": random.choice(DEGREES["Bachelor"]),
            "university": university,
            "start_date": f"{bachelors_start}-08-01",
            "end_date": f"{bachelors_end}-05-01",
            "gpa": round(random.uniform(3.2, 4.0), 2)
        })
        
        # Maybe Master's degree
        if random.random() > 0.5:
            masters_start = bachelors_end + random.randint(0, 2)
            masters_end = masters_start + 2
            
            # Fraudulent: overlapping with work or unrealistic timing
            if is_fraudulent and random.random() > 0.6:
                masters_start = bachelors_end - 1  # Overlap with bachelor's
            
            university = random.choice(FAKE_UNIVERSITIES if is_fraudulent and random.random() > 0.6
                                      else GENUINE_UNIVERSITIES)
            
            education.append({
                "degree": random.choice(DEGREES["Master"]),
                "university": university,
                "start_date": f"{masters_start}-08-01",
                "end_date": f"{masters_end}-05-01",
                "gpa": round(random.uniform(3.5, 4.0), 2)
            })
        
        return education
    
    def generate_experience(self, education: List[Dict], is_fraudulent: bool) -> List[Dict]:
        """Generate work experience"""
        experiences = []
        
        # Get graduation year
        grad_year = int(education[0]["end_date"].split("-")[0])
        current_year = datetime.now().year
        
        years_of_experience = current_year - grad_year
        num_jobs = min(max(1, years_of_experience // 2), 5)
        
        current_date = grad_year
        career_level = "entry"
        
        for i in range(num_jobs):
            # Duration in months
            duration_months = random.randint(12, 48)
            
            # Fraudulent: very short durations or overlaps
            if is_fraudulent and random.random() > 0.7:
                duration_months = random.randint(2, 8)  # Unrealistically short
            
            start_date = current_date
            end_date = start_date + (duration_months / 12)
            
            # Fraudulent: overlapping employment
            if is_fraudulent and random.random() > 0.8 and i > 0:
                start_date = current_date - 1  # Overlap with previous job
            
            # Determine job level
            years_exp = start_date - grad_year
            if years_exp >= 10:
                career_level = "executive" if is_fraudulent and random.random() > 0.7 else "senior"
            elif years_exp >= 5:
                career_level = "senior" if random.random() > 0.5 else "mid"
            elif years_exp >= 2:
                career_level = "mid"
            else:
                career_level = "entry"
            
            # Fraudulent: inflated titles
            if is_fraudulent and random.random() > 0.6:
                title = random.choice(INFLATED_TITLES)
            else:
                title = random.choice(JOB_TITLES_PROGRESSION[career_level])
            
            company = random.choice(FAKE_COMPANIES if is_fraudulent and random.random() > 0.5
                                   else GENUINE_COMPANIES)
            
            responsibilities = self.generate_responsibilities(career_level, is_fraudulent)
            
            experiences.append({
                "title": title,
                "company": company,
                "start_date": f"{int(start_date)}-{random.randint(1, 12):02d}-01",
                "end_date": f"{int(end_date)}-{random.randint(1, 12):02d}-01" if end_date < current_year else "Present",
                "responsibilities": responsibilities
            })
            
            current_date = end_date
            if current_date >= current_year:
                break
        
        return experiences
    
    def generate_responsibilities(self, level: str, is_fraudulent: bool) -> List[str]:
        """Generate job responsibilities"""
        responsibilities = []
        num_responsibilities = random.randint(3, 5)
        
        templates = {
            "entry": [
                "Developed and maintained software applications",
                "Collaborated with team members on project deliverables",
                "Participated in code reviews and testing",
                "Assisted in debugging and troubleshooting"
            ],
            "mid": [
                "Led development of key features for major products",
                "Mentored junior developers and conducted code reviews",
                "Designed and implemented scalable architecture",
                "Collaborated with cross-functional teams"
            ],
            "senior": [
                "Architected enterprise-level solutions serving millions of users",
                "Led team of 10+ engineers across multiple projects",
                "Established best practices and coding standards",
                "Drove technical strategy and innovation"
            ],
            "executive": [
                "Defined company-wide technical vision and strategy",
                "Managed engineering organization of 100+ people",
                "Drove digital transformation initiatives",
                "Built partnerships with major technology companies"
            ]
        }
        
        base_responsibilities = templates.get(level, templates["entry"])
        
        for _ in range(num_responsibilities):
            resp = random.choice(base_responsibilities)
            
            # Fraudulent: add exaggeration
            if is_fraudulent and random.random() > 0.6:
                exaggerators = ["single-handedly", "revolutionized", "transformed",
                              "best-in-class", "world-leading", "groundbreaking"]
                resp = f"{random.choice(exaggerators).capitalize()} {resp.lower()}"
            
            responsibilities.append(resp)
        
        return responsibilities
    
    def generate_skills(self, is_fraudulent: bool) -> List[str]:
        """Generate skills list"""
        skills = []
        
        # Select from different categories
        for category in ["programming", "web", "data", "cloud", "tools"]:
            num_skills = random.randint(1, 3)
            skills.extend(random.sample(SKILLS[category], min(num_skills, len(SKILLS[category]))))
        
        # Fraudulent: unrealistic number of skills or fake combinations
        if is_fraudulent and random.random() > 0.6:
            skills.extend(random.choice(UNREALISTIC_SKILL_COMBOS))
        
        # Fraudulent: add ML skills without relevant experience
        if is_fraudulent and random.random() > 0.7:
            skills.extend(random.sample(SKILLS["ml"], 3))
        
        return skills
    
    def generate_certifications(self, is_fraudulent: bool) -> List[str]:
        """Generate certifications"""
        certifications = []
        
        num_certs = random.randint(0, 3)
        
        for _ in range(num_certs):
            if is_fraudulent and random.random() > 0.5:
                certifications.append(random.choice(CERTIFICATIONS_FAKE))
            else:
                certifications.append(random.choice(CERTIFICATIONS_GENUINE))
        
        return certifications
    
    def create_resume_data(self, is_fraudulent: bool) -> Dict:
        """Create complete resume data structure"""
        contact = self.generate_contact()
        education = self.generate_education(is_fraudulent)
        experience = self.generate_experience(education, is_fraudulent)
        skills = self.generate_skills(is_fraudulent)
        certifications = self.generate_certifications(is_fraudulent)
        
        return {
            "contact": contact,
            "summary": self.generate_summary(experience, is_fraudulent),
            "education": education,
            "experience": experience,
            "skills": skills,
            "certifications": certifications,
            "is_fraudulent": is_fraudulent
        }
    
    def generate_summary(self, experience: List[Dict], is_fraudulent: bool) -> str:
        """Generate professional summary"""
        years = len(experience) * 2
        
        if is_fraudulent and random.random() > 0.5:
            return (f"Highly accomplished technology leader with {years}+ years of "
                   f"groundbreaking experience. Recognized as industry's top expert and "
                   f"best-in-class professional. Single-handedly transformed multiple "
                   f"Fortune 500 companies.")
        else:
            return (f"Experienced software engineer with {years} years in technology. "
                   f"Strong background in software development, team collaboration, "
                   f"and delivering quality solutions.")
    
    def generate_pdf(self, resume_data: Dict, filename: str):
        """Generate PDF resume"""
        pdf_path = self.output_dir / filename
        doc = SimpleDocTemplate(str(pdf_path), pagesize=letter)
        story = []
        styles = getSampleStyleSheet()
        
        # Title
        title_style = ParagraphStyle(
            'CustomTitle',
            parent=styles['Heading1'],
            fontSize=24,
            textColor=colors.HexColor('#2C3E50'),
            spaceAfter=30,
            alignment=1
        )
        
        story.append(Paragraph(resume_data['contact']['name'], title_style))
        
        # Contact info
        contact_text = f"{resume_data['contact']['email']} | {resume_data['contact']['phone']}"
        story.append(Paragraph(contact_text, styles['Normal']))
        story.append(Spacer(1, 0.2*inch))
        
        # Summary
        story.append(Paragraph("<b>PROFESSIONAL SUMMARY</b>", styles['Heading2']))
        story.append(Paragraph(resume_data['summary'], styles['Normal']))
        story.append(Spacer(1, 0.2*inch))
        
        # Experience
        story.append(Paragraph("<b>WORK EXPERIENCE</b>", styles['Heading2']))
        for exp in resume_data['experience']:
            story.append(Paragraph(f"<b>{exp['title']}</b> - {exp['company']}", styles['Normal']))
            story.append(Paragraph(f"{exp['start_date']} to {exp['end_date']}", styles['Normal']))
            for resp in exp['responsibilities']:
                story.append(Paragraph(f"• {resp}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))
        
        # Education
        story.append(Paragraph("<b>EDUCATION</b>", styles['Heading2']))
        for edu in resume_data['education']:
            story.append(Paragraph(f"<b>{edu['degree']}</b>", styles['Normal']))
            story.append(Paragraph(f"{edu['university']} ({edu['start_date']} - {edu['end_date']})", styles['Normal']))
            story.append(Paragraph(f"GPA: {edu['gpa']}", styles['Normal']))
            story.append(Spacer(1, 0.1*inch))
        
        # Skills
        story.append(Paragraph("<b>SKILLS</b>", styles['Heading2']))
        story.append(Paragraph(", ".join(resume_data['skills']), styles['Normal']))
        
        # Certifications
        if resume_data['certifications']:
            story.append(Spacer(1, 0.1*inch))
            story.append(Paragraph("<b>CERTIFICATIONS</b>", styles['Heading2']))
            for cert in resume_data['certifications']:
                story.append(Paragraph(f"• {cert}", styles['Normal']))
        
        doc.build(story)
        return pdf_path
    
    def generate_docx(self, resume_data: Dict, filename: str):
        """Generate DOCX resume"""
        docx_path = self.output_dir / filename
        doc = Document()
        
        # Title
        title = doc.add_heading(resume_data['contact']['name'], 0)
        title.alignment = 1  # Center
        
        # Contact
        contact_para = doc.add_paragraph()
        contact_para.add_run(f"{resume_data['contact']['email']} | {resume_data['contact']['phone']}")
        contact_para.alignment = 1
        
        # Summary
        doc.add_heading('PROFESSIONAL SUMMARY', 1)
        doc.add_paragraph(resume_data['summary'])
        
        # Experience
        doc.add_heading('WORK EXPERIENCE', 1)
        for exp in resume_data['experience']:
            doc.add_heading(f"{exp['title']} - {exp['company']}", 2)
            doc.add_paragraph(f"{exp['start_date']} to {exp['end_date']}")
            for resp in exp['responsibilities']:
                doc.add_paragraph(resp, style='List Bullet')
        
        # Education
        doc.add_heading('EDUCATION', 1)
        for edu in resume_data['education']:
            doc.add_heading(edu['degree'], 2)
            doc.add_paragraph(f"{edu['university']} ({edu['start_date']} - {edu['end_date']})")
            doc.add_paragraph(f"GPA: {edu['gpa']}")
        
        # Skills
        doc.add_heading('SKILLS', 1)
        doc.add_paragraph(", ".join(resume_data['skills']))
        
        # Certifications
        if resume_data['certifications']:
            doc.add_heading('CERTIFICATIONS', 1)
            for cert in resume_data['certifications']:
                doc.add_paragraph(cert, style='List Bullet')
        
        doc.save(str(docx_path))
        return docx_path
    
    def generate_dataset(self, num_resumes: int = 10000, fraud_ratio: float = 0.5):
        """Generate complete dataset of resumes"""
        print(f"Generating {num_resumes} synthetic resumes...")
        
        num_fraudulent = int(num_resumes * fraud_ratio)
        num_genuine = num_resumes - num_fraudulent
        
        for i in tqdm(range(num_resumes), desc="Generating resumes"):
            is_fraudulent = i < num_fraudulent
            
            # Create resume data
            resume_data = self.create_resume_data(is_fraudulent)
            
            # Randomly choose format
            format_type = random.choice(['pdf', 'docx'])
            filename = f"resume_{i:05d}.{format_type}"
            
            # Generate file
            if format_type == 'pdf':
                file_path = self.generate_pdf(resume_data, filename)
            else:
                file_path = self.generate_docx(resume_data, filename)
            
            # Store metadata
            self.metadata.append({
                'filename': filename,
                'label': 1 if is_fraudulent else 0,
                'fraud_type': 'fraudulent' if is_fraudulent else 'genuine',
                'format': format_type
            })
        
        # Save metadata
        metadata_path = self.output_dir / 'metadata.json'
        with open(metadata_path, 'w') as f:
            json.dump(self.metadata, f, indent=2)
        
        print(f"\nDataset generation complete!")
        print(f"Total resumes: {num_resumes}")
        print(f"Genuine: {num_genuine}")
        print(f"Fraudulent: {num_fraudulent}")
        print(f"Metadata saved to: {metadata_path}")


def main():
    """Main function to generate dataset"""
    output_dir = config.SYNTHETIC_DATA_DIR
    generator = ResumeGenerator(output_dir)
    
    # Generate dataset
    generator.generate_dataset(
        num_resumes=config.DATASET_CONFIG['total_resumes'],
        fraud_ratio=config.DATASET_CONFIG['fraud_ratio']
    )


if __name__ == "__main__":
    main()
